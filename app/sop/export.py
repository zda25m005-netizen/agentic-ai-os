"""Export the SOP as TXT / DOCX / PDF — document content only (no AI/telemetry)."""

from __future__ import annotations

import io
import re

from app.sop.models import SopDoc


def _clean(content: str) -> str:
    # strip stray markdown artifacts so the exported document reads as prose
    text = re.sub(r"^#{1,6}\s*", "", content, flags=re.M)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*(?!\*)(.+?)\*", r"\1", text)
    return text.strip()


def to_txt(doc: SopDoc) -> bytes:
    return _clean(doc.content).encode("utf-8")


def to_docx(doc: SopDoc) -> bytes:
    import docx

    d = docx.Document()
    d.add_heading(doc.title or "Statement of Purpose", level=1)
    for para in _clean(doc.content).split("\n\n"):
        if para.strip():
            d.add_paragraph(para.strip())
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def to_pdf(doc: SopDoc) -> bytes:
    from app.exec.pdf import PdfDoc, Section, build_pdf

    return build_pdf(
        PdfDoc(
            title=doc.title or "Statement of Purpose", sections=[Section("", _clean(doc.content))]
        )
    )


def export(doc: SopDoc, fmt: str) -> tuple[bytes, str, str]:
    fmt = (fmt or "pdf").lower()
    slug = re.sub(r"[^a-z0-9]+", "-", (doc.title or "sop").lower()).strip("-") or "sop"
    if fmt == "txt":
        return to_txt(doc), "text/plain", f"{slug}.txt"
    if fmt == "docx":
        return (
            to_docx(doc),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            f"{slug}.docx",
        )
    return to_pdf(doc), "application/pdf", f"{slug}.pdf"
